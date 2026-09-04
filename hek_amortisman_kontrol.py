from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import threading
import tkinter as tk
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None
    load_workbook = None


APP_TITLE = "HEK / Amortisman Eklenti Kontrolü"
SUPPORTED_TYPES = [("Excel dosyaları", "*.xlsx *.xlsm"), ("Tüm dosyalar", "*.*")]

HEK_PROCESS_COL = 37  # AK
HEK_INVENTORY_COL = 8  # H
HEK_SUBNUMBER_COL = 9  # I
HEK_RATIO_COL = 16  # P
AMORT_INVENTORY_COL = 1  # A
AMORT_SUBNUMBER_COL = 2  # B

YELLOW = "#F9D84A"
PALE_YELLOW = "#FFF8D6"
BLUE = "#2867A8"
PALE_BLUE = "#EAF3FB"
DARK = "#252A31"
GREEN = "#DDF2E3"
RED = "#FCE1E1"
ORANGE = "#FDE9C8"
WHITE = "#FFFFFF"
GRID = "#CBD5E1"
MUTED = "#64748B"

IGNORED_PROCESS_HEADERS = {
    "process id", "processid", "süreç no", "süreç numarası", "surec no", "surec numarasi"
}


def xl_color(color: str) -> str:
    return color.lstrip("#")


def resource_path(relative: str) -> Path:
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base / relative


def clean_identifier(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def identifier_key(value) -> str | None:
    text = clean_identifier(value)
    return text.casefold() if text is not None else None


def subnumber_key(value) -> str | None:
    text = clean_identifier(value)
    if text is None:
        return None
    compact = text.replace(" ", "")
    if re.fullmatch(r"[+-]?\d+(?:\.0+)?", compact):
        try:
            return str(int(Decimal(compact)))
        except (InvalidOperation, ValueError):
            pass
    return compact.casefold()


def natural_sort_key(value: str):
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def parse_process_input(text: str):
    values = []
    seen = set()
    duplicates = 0

    for raw_line in text.replace("\r", "\n").split("\n"):
        cells = raw_line.split("\t")
        first_cell = cells[0].strip()
        if not first_cell:
            continue
        if len([cell for cell in cells if cell.strip()]) > 1:
            continue
        if " ".join(first_cell.split()).casefold() in IGNORED_PROCESS_HEADERS:
            continue
        key = identifier_key(first_cell)
        if key in seen:
            duplicates += 1
            continue
        seen.add(key)
        values.append((key, first_cell))

    return values, duplicates


def process_line_statuses(text: str):
    """Return one UI status per visible line without restricting ID format."""
    statuses = []
    seen = set()
    for raw_line in text.replace("\r", "\n").split("\n"):
        cells = raw_line.split("\t")
        first_cell = cells[0].strip()
        nonempty_cells = [cell for cell in cells if cell.strip()]
        normalized = " ".join(first_cell.split()).casefold()
        if not nonempty_cells:
            status = "empty"
        elif len(nonempty_cells) > 1 or not first_cell:
            status = "invalid"
        elif normalized in IGNORED_PROCESS_HEADERS:
            status = "header"
        else:
            key = identifier_key(first_cell)
            if key in seen:
                status = "duplicate"
            else:
                seen.add(key)
                status = "valid"
        statuses.append(status)
    return statuses


def ratio_info(value):
    if value is None or str(value).strip() == "":
        return "(boş)", False

    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(str(value))
        percent = number * 100 if abs(number) <= 1 else number
    else:
        text = str(value).strip().replace("\u00a0", "").replace(" ", "")
        text = text.replace("%", "").replace(",", ".")
        try:
            percent = Decimal(text)
        except InvalidOperation:
            return str(value).strip(), False

    display = f"%{percent.normalize():f}"
    return display, percent == Decimal("100")


def trim_trailing_empty(values):
    result = list(values)
    while result and result[-1] is None:
        result.pop()
    return result


def display_header(value, index: int) -> str:
    text = clean_identifier(value)
    return text or f"Sütun {get_column_letter(index + 1)}"


def ordered_unique(values):
    seen = set()
    result = []
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def unique_join(values, separator=" | "):
    cleaned = [str(value) for value in values if value not in (None, "")]
    return separator.join(ordered_unique(cleaned))


def read_header(path: Path, sheet_name: str, header_row: int):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        row = next(worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        return trim_trailing_empty(row)
    finally:
        workbook.close()


def collect_hek(path: Path, sheet_name: str, header_row: int, requested_processes):
    requested_keys = {key for key, _display in requested_processes}
    requested_display = {key: display for key, display in requested_processes}
    found_processes = set()
    entries = []
    invalid_rows = []

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
        ):
            process_value = row[HEK_PROCESS_COL - 1] if len(row) >= HEK_PROCESS_COL else None
            process_key = identifier_key(process_value)
            if process_key not in requested_keys:
                continue

            found_processes.add(process_key)
            inventory_value = row[HEK_INVENTORY_COL - 1] if len(row) >= HEK_INVENTORY_COL else None
            subnumber_value = row[HEK_SUBNUMBER_COL - 1] if len(row) >= HEK_SUBNUMBER_COL else None
            ratio_value = row[HEK_RATIO_COL - 1] if len(row) >= HEK_RATIO_COL else None
            inventory_key = identifier_key(inventory_value)
            sub_key = subnumber_key(subnumber_value)
            ratio_display, ratio_full = ratio_info(ratio_value)

            if inventory_key is None or sub_key is None:
                invalid_rows.append({
                    "process": requested_display[process_key],
                    "row": row_number,
                    "inventory": clean_identifier(inventory_value) or "(boş)",
                    "subnumber": clean_identifier(subnumber_value) or "(boş)",
                    "ratio": ratio_display,
                    "reason": "HEK satırında envanter veya eklenti numarası boş",
                })
                continue

            entries.append({
                "process_key": process_key,
                "process": requested_display[process_key],
                "inventory_key": inventory_key,
                "inventory": clean_identifier(inventory_value),
                "sub_key": sub_key,
                "subnumber": clean_identifier(subnumber_value),
                "ratio": ratio_display,
                "ratio_full": ratio_full,
                "row": row_number,
            })
    finally:
        workbook.close()

    missing_processes = [display for key, display in requested_processes if key not in found_processes]
    return entries, invalid_rows, missing_processes


def collect_amortisation(path: Path, sheet_name: str, header_row: int, inventory_keys):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        header_values = next(
            worksheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ()
        )
        headers = trim_trailing_empty(header_values)
        source_column_count = max(len(headers), 2)
        headers.extend([None] * (source_column_count - len(headers)))

        rows_by_inventory = defaultdict(list)
        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
        ):
            inventory_value = row[AMORT_INVENTORY_COL - 1] if row else None
            inventory_key_value = identifier_key(inventory_value)
            if inventory_key_value not in inventory_keys:
                continue
            padded = list(row[:source_column_count])
            padded.extend([None] * (source_column_count - len(padded)))
            rows_by_inventory[inventory_key_value].append({
                "values": padded,
                "inventory": clean_identifier(inventory_value),
                "subnumber": clean_identifier(row[AMORT_SUBNUMBER_COL - 1]) if len(row) >= 2 else None,
                "sub_key": subnumber_key(row[AMORT_SUBNUMBER_COL - 1]) if len(row) >= 2 else None,
                "row": row_number,
            })

        return headers, rows_by_inventory
    finally:
        workbook.close()


def build_analysis(processes, hek_entries, invalid_hek_rows, missing_processes, headers, amort_rows):
    entries_by_inventory = defaultdict(list)
    request_by_pair = defaultdict(list)
    inventory_display = {}

    for entry in hek_entries:
        entries_by_inventory[entry["inventory_key"]].append(entry)
        request_by_pair[(entry["inventory_key"], entry["sub_key"])].append(entry)
        inventory_display.setdefault(entry["inventory_key"], entry["inventory"])

    summary_rows = []
    detail_groups = []
    unmatched_rows = list(invalid_hek_rows)

    sorted_inventories = sorted(entries_by_inventory, key=lambda key: natural_sort_key(inventory_display[key]))
    for inventory_key_value in sorted_inventories:
        hek_for_inventory = entries_by_inventory[inventory_key_value]
        process_names = ordered_unique(entry["process"] for entry in hek_for_inventory)
        rows = list(amort_rows.get(inventory_key_value, []))
        rows.sort(key=lambda item: natural_sort_key(item["sub_key"] or ""))
        available_pairs = {(inventory_key_value, row["sub_key"]) for row in rows}

        for entry in hek_for_inventory:
            if (inventory_key_value, entry["sub_key"]) not in available_pairs:
                unmatched_rows.append({
                    "process": entry["process"],
                    "row": entry["row"],
                    "inventory": entry["inventory"],
                    "subnumber": entry["subnumber"],
                    "ratio": entry["ratio"],
                    "reason": "HEK kaydı Amortisman dosyasında bulunamadı",
                })

        addon_rows = [row for row in rows if row["sub_key"] not in (None, "0")]
        requested_addons = [
            row for row in addon_rows if request_by_pair.get((inventory_key_value, row["sub_key"]))
        ]
        missing_addons = [
            row for row in addon_rows if not request_by_pair.get((inventory_key_value, row["sub_key"]))
        ]
        completion = len(requested_addons) / len(addon_rows) if addon_rows else 1.0
        partial_entries = [entry for entry in hek_for_inventory if not entry["ratio_full"]]

        problems = []
        if not rows:
            problems.append("Envanter Amortisman dosyasında bulunamadı")
        if missing_addons:
            problems.append(f"{len(missing_addons)} kalan eklenti var")
        if partial_entries:
            problems.append("Kısmi Çıkış Oranı (txtKismiCikisOrani) %100 değil")
        status = "Kontrol gerekli" if problems else "Tamam"

        summary_rows.append({
            "processes": "\n".join(process_names),
            "inventory": inventory_display[inventory_key_value],
            "addon_total": len(addon_rows),
            "requested_count": len(requested_addons),
            "requested_list": unique_join(row["subnumber"] for row in requested_addons) or "—",
            "missing_count": len(missing_addons),
            "missing_list": unique_join(row["subnumber"] for row in missing_addons) or "—",
            "completion": completion,
            "p_check": unique_join(
                f'{entry["process"]}: {entry["ratio"]}' for entry in partial_entries
            ) or "%100",
            "status": status,
            "explanation": "; ".join(problems) or "Bütün eklentiler hurdalanmak istenmiş",
        })

        detail_rows = []
        for row in rows:
            requests = request_by_pair.get((inventory_key_value, row["sub_key"]), [])
            is_main = row["sub_key"] == "0"
            if requests:
                requesting_processes = ordered_unique(entry["process"] for entry in requests)
                ratios = ordered_unique(entry["ratio"] for entry in requests)
                partial = any(not entry["ratio_full"] for entry in requests)
                if is_main:
                    status_text = "Ana ürün - hurdalanmak isteniyor"
                else:
                    status_text = "Hurdalanmak istenen eklenti"
                explanation = (
                    "Kısmi Çıkış Oranı (txtKismiCikisOrani) %100 değil"
                    if partial else "HEK Formunda mevcut"
                )
                kind = "partial" if partial else "requested"
            else:
                requesting_processes = []
                ratios = []
                if is_main:
                    status_text = "Ana ürün HEK Formunda yok"
                    explanation = "Ana ürün satırı hurdalama talebinde bulunamadı"
                    kind = "warning"
                else:
                    status_text = "Kalan eklenti"
                    explanation = "Hurdalanmak istenmemiş / eksik eklenti"
                    kind = "missing"

            detail_rows.append({
                **row,
                "request_status": status_text,
                "request_processes": "\n".join(requesting_processes) or "—",
                "ratios": unique_join(ratios) or "—",
                "explanation": explanation,
                "kind": kind,
            })

        detail_groups.append({
            "processes": "\n".join(process_names),
            "inventory": inventory_display[inventory_key_value],
            "rows": detail_rows,
        })

    for process in missing_processes:
        unmatched_rows.append({
            "process": process,
            "row": "—",
            "inventory": "—",
            "subnumber": "—",
            "ratio": "—",
            "reason": "Süreç numarası HEK Formunda bulunamadı",
        })

    return {
        "headers": headers,
        "summary": summary_rows,
        "detail_groups": detail_groups,
        "unmatched": unmatched_rows,
        "process_count": len(processes),
    }


def style_header(row_cells, fill_color=BLUE):
    for cell in row_cells:
        cell.fill = PatternFill("solid", fgColor=xl_color(fill_color))
        cell.font = Font(color=xl_color(WHITE), bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_result_workbook(result, output_path: Path):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Kontrol Özeti"
    detail = workbook.create_sheet("Filtrelenmiş Amortisman")
    unmatched = workbook.create_sheet("Eşleşmeyen Kayıtlar")

    thin_gray = Side(style="thin", color="D4DAE1")
    section_border = Border(bottom=Side(style="medium", color=xl_color(BLUE)))

    summary_headers = [
        "Süreç Numaraları", "Envanter Numarası", "Toplam Eklenti",
        "Hurdalanacak Eklenti", "Hurdalanacak Eklenti No", "Kalan Eklenti",
        "Kalan Eklenti No", "Eklenti Tamamlanma",
        "Kısmi Çıkış Oranı Kontrolü (txtKismiCikisOrani)", "Durum", "Açıklama",
    ]
    summary.append(summary_headers)
    style_header(summary[1])
    summary.row_dimensions[1].height = 34

    for item in result["summary"]:
        summary.append([
            item["processes"], item["inventory"], item["addon_total"],
            item["requested_count"], item["requested_list"], item["missing_count"],
            item["missing_list"], item["completion"], item["p_check"], item["status"],
            item["explanation"],
        ])
        row = summary.max_row
        summary.cell(row, 8).number_format = "0%"
        fill = GREEN if item["status"] == "Tamam" else ORANGE
        if item["missing_count"]:
            fill = RED
        for cell in summary[row]:
            cell.fill = PatternFill("solid", fgColor=xl_color(fill))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)

    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    summary.sheet_view.showGridLines = False
    summary_widths = [25, 19, 14, 18, 25, 14, 25, 18, 28, 18, 42]
    for index, width in enumerate(summary_widths, start=1):
        summary.column_dimensions[get_column_letter(index)].width = width

    source_headers = [display_header(value, index) for index, value in enumerate(result["headers"])]
    detail_headers = ["Süreç Numaraları"] + source_headers + [
        "HEK Talep Durumu", "Talep Eden Süreç(ler)",
        "Kısmi Çıkış Oranı (txtKismiCikisOrani)", "Kontrol Açıklaması"
    ]
    detail.append(detail_headers)
    style_header(detail[1])
    detail.row_dimensions[1].height = 38
    detail.freeze_panes = "B2"
    detail.sheet_view.showGridLines = False

    for group_number, group in enumerate(result["detail_groups"]):
        start_row = detail.max_row + 1
        for item in group["rows"]:
            detail.append([
                group["processes"], *item["values"], item["request_status"],
                item["request_processes"], item["ratios"], item["explanation"],
            ])
            current_row = detail.max_row
            if item["kind"] == "missing":
                fill_color = RED
            elif item["kind"] in {"partial", "warning"}:
                fill_color = ORANGE
            else:
                fill_color = GREEN if group_number % 2 == 0 else PALE_BLUE
            for cell in detail[current_row]:
                cell.fill = PatternFill("solid", fgColor=xl_color(fill_color))
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            detail.cell(current_row, 1).border = section_border

        end_row = detail.max_row
        if end_row >= start_row:
            if end_row > start_row:
                detail.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            process_cell = detail.cell(start_row, 1)
            process_cell.value = group["processes"]
            process_cell.font = Font(bold=True, color=xl_color(DARK))
            process_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            process_cell.fill = PatternFill("solid", fgColor=xl_color(YELLOW))
            process_cell.border = Border(
                left=Side(style="medium", color=xl_color(BLUE)),
                right=Side(style="medium", color=xl_color(BLUE)),
                top=Side(style="medium", color=xl_color(BLUE)),
                bottom=Side(style="medium", color=xl_color(BLUE)),
            )

    if detail.max_row > 1:
        detail.auto_filter.ref = f"B1:{get_column_letter(detail.max_column)}{detail.max_row}"
    detail.column_dimensions["A"].width = 25
    for index, header in enumerate(source_headers, start=2):
        folded = header.casefold()
        if any(word in folded for word in ("description", "tanım", "tanim", "açıklama", "aciklama")):
            width = 42
        elif any(word in folded for word in ("date", "tarih")):
            width = 14
        else:
            width = min(max(len(header) + 2, 12), 24)
        detail.column_dimensions[get_column_letter(index)].width = width
    for index, width in enumerate([28, 24, 22, 42], start=2 + len(source_headers)):
        detail.column_dimensions[get_column_letter(index)].width = width

    unmatched_headers = [
        "Süreç Numarası", "HEK Satırı", "Envanter Numarası", "Eklenti Numarası",
        "Kısmi Çıkış Oranı (txtKismiCikisOrani)", "Açıklama",
    ]
    unmatched.append(unmatched_headers)
    style_header(unmatched[1])
    for item in result["unmatched"]:
        unmatched.append([
            item["process"], item["row"], item["inventory"], item["subnumber"],
            item["ratio"], item["reason"],
        ])
        for cell in unmatched[unmatched.max_row]:
            cell.fill = PatternFill("solid", fgColor=xl_color(RED))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    unmatched.freeze_panes = "A2"
    unmatched.auto_filter.ref = unmatched.dimensions
    unmatched.sheet_view.showGridLines = False
    for index, width in enumerate([24, 12, 20, 18, 22, 48], start=1):
        unmatched.column_dimensions[get_column_letter(index)].width = width

    for sheet in (summary, detail, unmatched):
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.outlinePr.summaryBelow = True

    workbook.save(output_path)


def run_analysis(process_text, hek_config, amort_config, output_path: Path):
    processes, duplicates = parse_process_input(process_text)
    if not processes:
        raise ValueError("En az bir süreç numarası yapıştırın.")

    hek_entries, invalid_rows, missing_processes = collect_hek(
        hek_config["path"], hek_config["sheet"], hek_config["header_row"], processes
    )
    inventory_keys = {entry["inventory_key"] for entry in hek_entries}
    headers, amort_rows = collect_amortisation(
        amort_config["path"], amort_config["sheet"], amort_config["header_row"], inventory_keys
    )
    result = build_analysis(
        processes, hek_entries, invalid_rows, missing_processes, headers, amort_rows
    )
    result["duplicates"] = duplicates
    result["hek_entry_count"] = len(hek_entries)
    result["inventory_count"] = len(inventory_keys)
    result["missing_addon_count"] = sum(item["missing_count"] for item in result["summary"])
    result["partial_inventory_count"] = sum(
        1 for item in result["summary"] if item["p_check"] != "%100"
    )
    write_result_workbook(result, output_path)
    return result


class ProcessInput(tk.Frame):
    """Excel-like single-column editor with synchronized row numbers."""

    def __init__(self, master, on_change):
        super().__init__(master, bg=BLUE, padx=2, pady=2)
        self.on_change = on_change
        self.columnconfigure(1, weight=1)
        self.rowconfigure(1, weight=1)

        tk.Label(
            self, text="SATIR", bg=BLUE, fg=WHITE, font=("TkDefaultFont", 9, "bold"),
            width=6, pady=5,
        ).grid(row=0, column=0, sticky="ew", padx=(0, 1))
        tk.Label(
            self, text="SÜREÇ NUMARASI — Excel'den tek sütun yapıştırın", bg=BLUE, fg=WHITE,
            font=("TkDefaultFont", 9, "bold"), anchor="w", padx=10, pady=5,
        ).grid(row=0, column=1, sticky="ew")

        self.line_numbers = tk.Text(
            self, width=6, height=6, bg="#EEF2F6", fg=MUTED, relief="flat",
            borderwidth=0, padx=4, pady=6, takefocus=0, cursor="arrow",
            font=("TkFixedFont", 11), state="disabled", wrap="none",
        )
        self.line_numbers.grid(row=1, column=0, sticky="nsew", padx=(0, 1))

        self.text = tk.Text(
            self, height=6, bg=WHITE, fg=DARK, insertbackground=DARK, relief="flat",
            borderwidth=0, padx=10, pady=6, undo=True, wrap="none", font=("TkFixedFont", 11),
        )
        self.text.grid(row=1, column=1, sticky="nsew")
        scrollbar = ttk.Scrollbar(self, orient="vertical", command=self._scroll_both)
        scrollbar.grid(row=1, column=2, sticky="ns")
        self.text.configure(yscrollcommand=lambda first, last: self._on_text_scroll(scrollbar, first, last))

        self.text.tag_configure("duplicate", background="#FFF1A8")
        self.text.tag_configure("invalid", background="#FFD6D6", foreground="#9B1C1C")
        self.text.tag_configure("header", background=PALE_BLUE, foreground=BLUE)
        self.text.bind("<<Modified>>", self._content_changed)
        self.text.bind("<MouseWheel>", lambda _event: self.after_idle(self._sync_gutter))
        self.text.bind("<Button-4>", lambda _event: self.after_idle(self._sync_gutter))
        self.text.bind("<Button-5>", lambda _event: self.after_idle(self._sync_gutter))
        self.text.edit_modified(False)
        self.refresh()

    def _scroll_both(self, *args):
        self.text.yview(*args)
        self.line_numbers.yview(*args)

    def _on_text_scroll(self, scrollbar, first, last):
        scrollbar.set(first, last)
        self.line_numbers.yview_moveto(first)

    def _sync_gutter(self):
        self.line_numbers.yview_moveto(self.text.yview()[0])

    def _content_changed(self, _event=None):
        if not self.text.edit_modified():
            return
        self.text.edit_modified(False)
        self.refresh()
        self.on_change()

    def refresh(self):
        content = self.get()
        statuses = process_line_statuses(content)
        visible_count = max(1, len(content.split("\n")))
        numbers = "\n".join(str(number) for number in range(1, visible_count + 1))

        self.line_numbers.configure(state="normal")
        self.line_numbers.delete("1.0", "end")
        self.line_numbers.insert("1.0", numbers)
        self.line_numbers.configure(state="disabled")

        for tag in ("duplicate", "invalid", "header"):
            self.text.tag_remove(tag, "1.0", "end")
        for line_number, status in enumerate(statuses, start=1):
            if status in {"duplicate", "invalid", "header"}:
                self.text.tag_add(status, f"{line_number}.0", f"{line_number}.end")
        self.after_idle(self._sync_gutter)

    def get(self):
        return self.text.get("1.0", "end-1c")

    def replace(self, value: str):
        self.text.delete("1.0", "end")
        self.text.insert("1.0", value)
        self.text.edit_modified(False)
        self.refresh()
        self.on_change()

    def clear(self):
        self.replace("")


class ExcelCard(ttk.LabelFrame):
    def __init__(self, master, title, kind):
        super().__init__(master, text=title, padding=12, style="Card.TLabelframe")
        self.kind = kind
        self.path: Path | None = None

        self.columnconfigure(1, weight=1)
        ttk.Label(self, text="Dosya", style="Card.TLabel").grid(row=0, column=0, sticky="w", padx=(0, 8))
        self.file_var = tk.StringVar(value="Dosya seçilmedi")
        ttk.Entry(self, textvariable=self.file_var, state="readonly").grid(row=0, column=1, sticky="ew")
        ttk.Button(self, text="Dosya Seç", command=self.choose_file, style="Blue.TButton").grid(
            row=0, column=2, padx=(8, 0)
        )

        ttk.Label(self, text="Sayfa", style="Card.TLabel").grid(row=1, column=0, sticky="w", pady=(8, 0))
        self.sheet_var = tk.StringVar()
        self.sheet_box = ttk.Combobox(self, textvariable=self.sheet_var, state="readonly")
        self.sheet_box.grid(row=1, column=1, sticky="ew", pady=(8, 0))
        self.sheet_box.bind("<<ComboboxSelected>>", lambda _event: self.validate_columns())

        ttk.Label(self, text="Başlık satırı", style="Card.TLabel").grid(
            row=1, column=2, sticky="e", padx=(8, 76), pady=(8, 0)
        )
        self.header_var = tk.IntVar(value=1)
        ttk.Spinbox(self, from_=1, to=1000, width=6, textvariable=self.header_var,
                    command=self.validate_columns).grid(row=1, column=2, sticky="e", pady=(8, 0))

        self.status_var = tk.StringVar(value="Dosya bekleniyor")
        ttk.Label(self, textvariable=self.status_var, style="Hint.TLabel", wraplength=780).grid(
            row=2, column=0, columnspan=3, sticky="w", pady=(9, 0)
        )

    def choose_file(self):
        filename = filedialog.askopenfilename(title=f"{self['text']} seç", filetypes=SUPPORTED_TYPES)
        if not filename:
            return
        path = Path(filename)
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            messagebox.showerror("Desteklenmeyen dosya", "Lütfen .xlsx veya .xlsm dosyası seçin.")
            return
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()
        except Exception as exc:
            messagebox.showerror("Dosya açılamadı", f"Excel dosyası okunamadı:\n{exc}")
            return
        self.path = path
        self.file_var.set(str(path))
        self.sheet_box["values"] = sheet_names
        self.sheet_var.set(sheet_names[0])
        self.validate_columns()

    def validate_columns(self):
        if not self.path or not self.sheet_var.get():
            return False
        try:
            header_row = int(self.header_var.get())
            headers = read_header(self.path, self.sheet_var.get(), header_row)
        except Exception as exc:
            self.status_var.set(f"⚠ Başlıklar okunamadı: {exc}")
            return False

        if self.kind == "hek":
            required = [
                (HEK_PROCESS_COL, "AK / Process ID"),
                (HEK_INVENTORY_COL, "H / Envanter"),
                (HEK_SUBNUMBER_COL, "I / Eklenti"),
                (HEK_RATIO_COL, "Kısmi Çıkış Oranı / txtKismiCikisOrani"),
            ]
        else:
            required = [
                (AMORT_INVENTORY_COL, "A / Immobilisation"),
                (AMORT_SUBNUMBER_COL, "B / Numero subsidiaire"),
            ]

        missing = [label for index, label in required if len(headers) < index]
        if missing:
            self.status_var.set("⚠ Eksik sütun: " + ", ".join(missing))
            return False

        detected = []
        for index, label in required:
            detected.append(f"{label}: {display_header(headers[index - 1], index - 1)}")
        self.status_var.set("✓ " + "   |   ".join(detected))
        return True

    def config(self):
        if not self.path:
            raise ValueError(f"{self['text']} seçilmedi.")
        if not self.validate_columns():
            raise ValueError(f"{self['text']} sütunları doğrulanamadı.")
        return {
            "path": self.path,
            "sheet": self.sheet_var.get(),
            "header_row": int(self.header_var.get()),
        }


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1120x820")
        self.minsize(900, 720)
        self.configure(bg=PALE_YELLOW)
        self.last_output: Path | None = None
        self._configure_styles()

        shell = ttk.Frame(self, padding=18, style="Shell.TFrame")
        shell.pack(fill="both", expand=True)
        shell.columnconfigure(0, weight=1)
        shell.rowconfigure(4, weight=1)

        header = ttk.Frame(shell, style="Header.TFrame", padding=14)
        header.grid(row=0, column=0, sticky="ew", pady=(0, 12))
        header.columnconfigure(1, weight=1)
        self._load_logo(header)
        ttk.Label(header, text="HEK / Amortisman Eklenti Kontrolü", style="Title.TLabel").grid(
            row=0, column=1, sticky="w"
        )
        ttk.Label(
            header,
            text="Süreçleri bulur, envanter eklentilerini kontrol eder ve düzenli Excel raporu üretir.",
            style="Subtitle.TLabel",
        ).grid(row=1, column=1, sticky="w", pady=(4, 0))
        ttk.Button(header, text="?  YARDIM", command=self.show_help, style="Help.TButton").grid(
            row=0, column=2, rowspan=2, sticky="e", padx=(16, 0)
        )

        process_card = ttk.LabelFrame(shell, text="1  Süreç Numaraları", padding=12, style="Card.TLabelframe")
        process_card.grid(row=1, column=0, sticky="ew", pady=(0, 10))
        process_card.columnconfigure(0, weight=1)
        self.process_input = ProcessInput(process_card, self.update_process_count)
        self.process_input.grid(row=0, column=0, columnspan=3, sticky="ew")
        ttk.Button(process_card, text="Panodan Yapıştır", command=self.paste_processes,
                   style="Yellow.TButton").grid(row=1, column=0, sticky="w", pady=(8, 0))
        ttk.Button(process_card, text="Temizle", command=self.clear_processes).grid(
            row=1, column=1, sticky="w", padx=8, pady=(8, 0)
        )
        self.process_count_var = tk.StringVar(value="Süreç numaralarını Excel'den kopyalayıp buraya yapıştırın")
        ttk.Label(process_card, textvariable=self.process_count_var, style="Hint.TLabel").grid(
            row=1, column=2, sticky="e", pady=(8, 0)
        )

        cards = ttk.Frame(shell, style="Shell.TFrame")
        cards.grid(row=2, column=0, sticky="ew")
        cards.columnconfigure(0, weight=1)
        cards.columnconfigure(1, weight=1)
        self.hek_card = ExcelCard(cards, "2  HEK Formu", "hek")
        self.amort_card = ExcelCard(cards, "3  Amortisman Dosyası", "amort")
        self.hek_card.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
        self.amort_card.grid(row=0, column=1, sticky="nsew", padx=(6, 0))

        action = ttk.Frame(shell, style="Shell.TFrame")
        action.grid(row=3, column=0, sticky="ew", pady=12)
        self.run_button = ttk.Button(action, text="ANALİZİ BAŞLAT", command=self.start_analysis,
                                     style="Primary.TButton")
        self.run_button.pack(side="left")
        self.open_button = ttk.Button(action, text="Sonuç Excel'ini Aç", command=self.open_output,
                                      state="disabled")
        self.open_button.pack(side="left", padx=8)
        self.folder_button = ttk.Button(action, text="Klasörde Göster", command=self.show_output_folder,
                                        state="disabled")
        self.folder_button.pack(side="left")

        self.progress = ttk.Progressbar(action, mode="indeterminate", length=210)
        self.progress.pack(side="right")

        result_card = ttk.LabelFrame(shell, text="Sonuç", padding=16, style="Card.TLabelframe")
        result_card.grid(row=4, column=0, sticky="nsew")
        result_card.columnconfigure(0, weight=1)
        result_card.rowconfigure(1, weight=1)
        self.status_var = tk.StringVar(value="Hazır")
        ttk.Label(result_card, textvariable=self.status_var, style="ResultTitle.TLabel").grid(
            row=0, column=0, sticky="w"
        )
        self.result_var = tk.StringVar(value="Üç girişi tamamlayıp analizi başlatın.")
        ttk.Label(result_card, textvariable=self.result_var, style="Result.TLabel",
                  justify="left", wraplength=980).grid(row=1, column=0, sticky="nw", pady=(12, 0))

    def _configure_styles(self):
        style = ttk.Style(self)
        if "clam" in style.theme_names():
            style.theme_use("clam")
        style.configure("Shell.TFrame", background=PALE_YELLOW)
        style.configure("Header.TFrame", background=BLUE)
        style.configure("Title.TLabel", background=BLUE, foreground=WHITE, font=("TkDefaultFont", 19, "bold"))
        style.configure("Subtitle.TLabel", background=BLUE, foreground="#EAF3FB", font=("TkDefaultFont", 10))
        style.configure("Card.TLabelframe", background=WHITE, bordercolor=BLUE, relief="solid")
        style.configure("Card.TLabelframe.Label", background=YELLOW, foreground=DARK,
                        font=("TkDefaultFont", 11, "bold"))
        style.configure("Card.TLabel", background=WHITE, foreground=DARK)
        style.configure("Hint.TLabel", background=WHITE, foreground="#4E5965")
        style.configure("ResultTitle.TLabel", background=WHITE, foreground=BLUE,
                        font=("TkDefaultFont", 15, "bold"))
        style.configure("Result.TLabel", background=WHITE, foreground=DARK, font=("TkDefaultFont", 11))
        style.configure("HelpTitle.TLabel", background=PALE_YELLOW, foreground=BLUE,
                        font=("TkDefaultFont", 18, "bold"))
        style.configure("HelpText.TLabel", background=PALE_YELLOW, foreground=DARK,
                        font=("TkDefaultFont", 10))
        style.configure("HelpCard.TFrame", background=WHITE, relief="solid", borderwidth=1)
        style.configure("HelpSection.TLabel", background=WHITE, foreground=BLUE,
                        font=("TkDefaultFont", 11, "bold"))
        style.configure("HelpCardText.TLabel", background=WHITE, foreground=DARK,
                        font=("TkDefaultFont", 10))
        style.configure("Primary.TButton", background=BLUE, foreground=WHITE,
                        font=("TkDefaultFont", 12, "bold"), padding=(18, 10))
        style.map("Primary.TButton", background=[("active", "#1F527F")])
        style.configure("Yellow.TButton", background=YELLOW, foreground=DARK, padding=(12, 6))
        style.map("Yellow.TButton", background=[("active", "#E8C62E")])
        style.configure("Blue.TButton", background=BLUE, foreground=WHITE)
        style.map("Blue.TButton", background=[("active", "#1F527F")])
        style.configure("Help.TButton", background=WHITE, foreground=BLUE,
                        font=("TkDefaultFont", 10, "bold"), padding=(12, 7))
        style.map("Help.TButton", background=[("active", PALE_BLUE)])

    def _load_logo(self, parent):
        try:
            header_image = tk.PhotoImage(file=str(resource_path("assets/renault-logo-header.png")))
            self.icon_image = tk.PhotoImage(file=str(resource_path("assets/renault-logo.png")))
            self.logo_image = header_image.subsample(3, 3)
            tk.Label(parent, image=self.logo_image, bg=BLUE, borderwidth=0).grid(
                row=0, column=0, rowspan=2, padx=(0, 16)
            )
            self.iconphoto(True, self.icon_image)
        except Exception:
            ttk.Label(parent, text="◉", style="Title.TLabel", font=("TkDefaultFont", 30, "bold")).grid(
                row=0, column=0, rowspan=2, padx=(0, 16)
            )

    def paste_processes(self):
        try:
            value = self.clipboard_get()
        except tk.TclError:
            messagebox.showinfo("Pano boş", "Panoda yapıştırılabilecek metin bulunamadı.")
            return
        self.process_input.replace(value)

    def clear_processes(self):
        self.process_input.clear()

    def update_process_count(self):
        process_text = self.process_input.get()
        processes, duplicates = parse_process_input(process_text)
        invalid = process_line_statuses(process_text).count("invalid")
        if not processes:
            if invalid:
                self.process_count_var.set(f"⚠ {invalid} satırda birden fazla sütun var")
            else:
                self.process_count_var.set("Süreç numaralarını Excel'den tek sütun olarak yapıştırın")
        else:
            notes = [f"{len(processes)} benzersiz süreç"]
            if duplicates:
                notes.append(f"{duplicates} tekrar sarı işaretlendi")
            if invalid:
                notes.append(f"{invalid} çok sütunlu satır kırmızı işaretlendi")
            self.process_count_var.set(" • ".join(notes))

    def show_help(self):
        dialog = tk.Toplevel(self)
        dialog.title("Yardım — Sonuç Excel'i nasıl okunur?")
        dialog.geometry("720x610")
        dialog.minsize(620, 520)
        dialog.configure(bg=PALE_YELLOW)
        dialog.transient(self)
        dialog.grab_set()

        panel = ttk.Frame(dialog, padding=20, style="Shell.TFrame")
        panel.pack(fill="both", expand=True)
        panel.columnconfigure(0, weight=1)
        ttk.Label(panel, text="Sonuç Excel'i nasıl okunur?", style="HelpTitle.TLabel").grid(
            row=0, column=0, sticky="w"
        )
        ttk.Label(
            panel,
            text=(
                "Yardım dosyası, uygulamanın ürettiği üç sonuç sayfasını örnek verilerle gösterir. "
                "Açıklamalar sayfasında her sonuç sütununun ve satır türünün anlamı bulunur."
            ),
            style="HelpText.TLabel", wraplength=660, justify="left",
        ).grid(row=1, column=0, sticky="w", pady=(10, 16))

        sections = [
            (
                "Kontrol Özeti",
                "Her satır bir envanter grubunu gösterir. Hurdalanacak Eklenti talepte bulunanları, "
                "Kalan Eklenti ise Amortisman dosyasında olup hurdalama talebinde bulunmayanları gösterir.",
            ),
            (
                "Filtrelenmiş Amortisman",
                "Her satır bir ana ürün veya eklentidir. Amortisman dosyasının bütün kaynak sütunları "
                "korunur; sağ tarafta HEK talep durumu, talep eden süreç ve txtKismiCikisOrani yer alır.",
            ),
            (
                "Eşleşmeyen Kayıtlar",
                "HEK Formunda envanteri veya eklenti numarası eksik olan, Amortisman dosyasında "
                "bulunamayan ya da süreç numarası HEK Formunda bulunmayan kayıtları gösterir.",
            ),
            (
                "Renkler ve oran",
                "Yeşil satırlar uygun, turuncu satırlar txtKismiCikisOrani %100 olmadığı için kontrol "
                "gerektiren, kırmızı satırlar ise kalan veya eşleşmeyen kayıtları gösterir.",
            ),
        ]
        next_row = 2
        for title, body in sections:
            card = ttk.Frame(panel, padding=12, style="HelpCard.TFrame")
            card.grid(row=next_row, column=0, sticky="ew", pady=(0, 8))
            card.columnconfigure(0, weight=1)
            ttk.Label(card, text=title, style="HelpSection.TLabel").grid(row=0, column=0, sticky="w")
            ttk.Label(card, text=body, style="HelpCardText.TLabel", wraplength=625,
                      justify="left").grid(row=1, column=0, sticky="w", pady=(4, 0))
            next_row += 1

        buttons = ttk.Frame(panel, style="Shell.TFrame")
        buttons.grid(row=next_row, column=0, sticky="ew", pady=(10, 0))
        ttk.Button(buttons, text="ÖRNEK SONUÇ EXCEL'İNİ İNDİR", command=self.download_example_workbook,
                   style="Primary.TButton").pack(side="left")
        ttk.Button(buttons, text="Kapat", command=dialog.destroy).pack(side="right")

    def download_example_workbook(self):
        source = resource_path("assets/HEK_Amortisman_Ornek_Cikti.xlsx")
        if not source.exists():
            messagebox.showerror("Örnek dosya bulunamadı", "Örnek Excel dosyası uygulamaya eklenememiş.")
            return
        filename = filedialog.asksaveasfilename(
            title="Örnek Excel dosyasını kaydet",
            defaultextension=".xlsx",
            initialfile="HEK_Amortisman_Ornek_Cikti.xlsx",
            filetypes=[("Excel dosyası", "*.xlsx")],
        )
        if not filename:
            return
        try:
            shutil.copyfile(source, filename)
        except Exception as exc:
            messagebox.showerror("Dosya kaydedilemedi", f"Örnek Excel dosyası kaydedilemedi:\n{exc}")
            return
        messagebox.showinfo("Örnek dosya hazır", f"Örnek Excel dosyası kaydedildi:\n{filename}")

    def start_analysis(self):
        try:
            process_text = self.process_input.get()
            invalid_count = process_line_statuses(process_text).count("invalid")
            if invalid_count:
                raise ValueError(
                    f"{invalid_count} satırda birden fazla Excel sütunu var. "
                    "Yalnızca süreç numarası sütununu kopyalayın."
                )
            processes, _duplicates = parse_process_input(process_text)
            if not processes:
                raise ValueError("Süreç numaralarını Excel'den kopyalayıp yapıştırın.")
            hek_config = self.hek_card.config()
            amort_config = self.amort_card.config()
        except Exception as exc:
            messagebox.showerror("Eksik bilgi", str(exc))
            return

        filename = filedialog.asksaveasfilename(
            title="Analiz Excel dosyasını kaydet",
            defaultextension=".xlsx",
            initialfile="HEK_Amortisman_Eklenti_Kontrolu.xlsx",
            filetypes=[("Excel dosyası", "*.xlsx")],
        )
        if not filename:
            return

        self.run_button.configure(state="disabled")
        self.open_button.configure(state="disabled")
        self.folder_button.configure(state="disabled")
        self.progress.start(10)
        self.status_var.set("Analiz yapılıyor…")
        self.result_var.set("HEK Formu ve Amortisman dosyası birer kez taranıyor. Lütfen bekleyin.")

        thread = threading.Thread(
            target=self._analysis_worker,
            args=(process_text, hek_config, amort_config, Path(filename)),
            daemon=True,
        )
        thread.start()

    def _analysis_worker(self, process_text, hek_config, amort_config, output_path):
        try:
            result = run_analysis(process_text, hek_config, amort_config, output_path)
        except Exception as exc:
            self.after(0, lambda: self._analysis_failed(exc))
            return
        self.after(0, lambda: self._analysis_finished(result, output_path))

    def _analysis_failed(self, exc):
        self.progress.stop()
        self.run_button.configure(state="normal")
        self.status_var.set("Analiz tamamlanamadı")
        self.result_var.set(str(exc))
        messagebox.showerror("Analiz tamamlanamadı", str(exc))

    def _analysis_finished(self, result, output_path):
        self.progress.stop()
        self.run_button.configure(state="normal")
        self.open_button.configure(state="normal")
        self.folder_button.configure(state="normal")
        self.last_output = output_path
        matched_processes = result["process_count"] - len([
            row for row in result["unmatched"] if row["reason"] == "Süreç numarası HEK Formunda bulunamadı"
        ])
        self.status_var.set("Analiz tamamlandı")
        self.result_var.set(
            f"✓ {matched_processes}/{result['process_count']} süreç HEK Formunda bulundu\n"
            f"✓ {result['inventory_count']} envanter numarası analiz edildi\n"
            f"⚠ {result['missing_addon_count']} kalan eklenti bulundu\n"
            f"⚠ {result['partial_inventory_count']} envanter grubunda Kısmi Çıkış Oranı "
            f"(txtKismiCikisOrani) %100 değil\n"
            f"⚠ {len(result['unmatched'])} eşleşmeyen veya eksik kayıt var\n\n"
            f"Excel kaydedildi: {output_path}"
        )
        messagebox.showinfo("Tamamlandı", "Analiz Excel dosyası başarıyla oluşturuldu.")

    def open_output(self):
        if self.last_output:
            self._open_path(self.last_output)

    def show_output_folder(self):
        if not self.last_output:
            return
        if sys.platform == "win32":
            subprocess.Popen(["explorer", "/select,", str(self.last_output)])
        elif sys.platform == "darwin":
            subprocess.Popen(["open", "-R", str(self.last_output)])
        else:
            subprocess.Popen(["xdg-open", str(self.last_output.parent)])

    @staticmethod
    def _open_path(path: Path):
        if sys.platform == "win32":
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])


if __name__ == "__main__":
    if load_workbook is None:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror(
            "Eksik bileşen",
            "Excel okuyucu bulunamadı. Önce openpyxl paketini kurun.",
        )
        raise SystemExit(1)
    App().mainloop()
