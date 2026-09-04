import importlib.util
import sys
import tempfile
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook


PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT / "hek_amortisman_kontrol.py"
spec = importlib.util.spec_from_file_location("hek_app", SOURCE)
app = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = app
spec.loader.exec_module(app)


def make_hek(path):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("HEK")
    headers = [f"HEK Sütun {i}" for i in range(1, 38)]
    headers[7] = "Envanter Numarası"
    headers[8] = "Eklenti Numarası"
    headers[15] = "txtKismiCikisOrani"
    headers[36] = "Process ID"
    sheet.append(headers)

    def row(process, inventory, subnumber, ratio):
        values = [None] * 37
        values[7] = inventory
        values[8] = subnumber
        values[15] = ratio
        values[36] = process
        return values

    sheet.append(row("SUREC-001", 87593849, 0, "100%"))
    sheet.append(row("SUREC-001", 87593849, 1, 1))
    sheet.append(row("SUREC-002", 87593849, 2, "75%"))
    sheet.append(row("SUREC-002", 87593849, 9, "100%"))
    sheet.append(row("BASKA-SUREC", 99999999, 0, "100%"))
    workbook.save(path)


def make_amortisation(path, unrelated_rows=10000):
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Format")
    headers = [
        "Immobilisation", "Numero subsidiaire", "Tarih", "Açıklama",
        "Brüt Değer", "Amortisman", "Net Değer", "Para Birimi",
    ]
    headers.extend([f"Amort Sütun {i}" for i in range(9, 20)])
    headers.append("Val Compt Fin")
    sheet.append(headers)

    for subnumber, description, value in [
        (0, "Ana ürün", 1000),
        (1, "Birinci eklenti", 200),
        (2, "İkinci eklenti", 300),
        (3, "Üçüncü eklenti", 400),
    ]:
        values = [87593849, subnumber, "31.12.2025", description, value, -value, 0, "EUR"]
        values.extend([f"V{i}" for i in range(9, 20)])
        values.append(0)
        sheet.append(values)

    for index in range(unrelated_rows):
        values = [70000000 + index, 0, "31.12.2025", "İlgisiz kayıt", 1, -1, 0, "EUR"]
        values.extend([None] * 12)
        sheet.append(values)

    workbook.save(path)


def run_test():
    assert (PROJECT / "assets" / "HEK_Amortisman_Ornek_Cikti.xlsx").exists()
    temp_dir = Path(tempfile.mkdtemp(prefix="hek-amort-test-", dir="/private/tmp"))
    hek_path = temp_dir / "HEK_Formu.xlsx"
    amort_path = temp_dir / "Amortisman.xlsx"
    output_dir = PROJECT.parents[1] / "outputs"
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / "HEK_Amortisman_Ornek_Sonuc.xlsx"

    make_hek(hek_path)
    make_amortisation(amort_path)
    started = time.perf_counter()
    result = app.run_analysis(
        "Process ID\nSUREC-001\nSUREC-002\nSUREC-002\nSUREC-404\n",
        {"path": hek_path, "sheet": "HEK", "header_row": 1},
        {"path": amort_path, "sheet": "Format", "header_row": 1},
        output_path,
    )
    elapsed = time.perf_counter() - started

    assert result["process_count"] == 3
    assert result["duplicates"] == 1
    parsed, duplicates = app.parse_process_input("SUREC-001\nSUREC-001\nSUREC-002\tYANLIS-SUTUN")
    assert [display for _key, display in parsed] == ["SUREC-001"]
    assert duplicates == 1
    assert app.process_line_statuses("SUREC-001\nSUREC-001\nSUREC-002\tYANLIS-SUTUN") == [
        "valid", "duplicate", "invalid"
    ]
    assert result["inventory_count"] == 1
    assert result["missing_addon_count"] == 1
    summary = result["summary"][0]
    assert summary["processes"] == "SUREC-001\nSUREC-002"
    assert summary["addon_total"] == 3
    assert summary["requested_count"] == 2
    assert summary["missing_list"] == "3"
    assert round(summary["completion"], 4) == 0.6667
    assert "SUREC-002: %75" in summary["p_check"]
    assert any(item["subnumber"] == "9" for item in result["unmatched"])
    assert any(item["process"] == "SUREC-404" for item in result["unmatched"])

    workbook = load_workbook(output_path, data_only=False)
    assert workbook.sheetnames == ["Kontrol Özeti", "Filtrelenmiş Amortisman", "Eşleşmeyen Kayıtlar"]
    detail = workbook["Filtrelenmiş Amortisman"]
    assert "A2:A5" in {str(item) for item in detail.merged_cells.ranges}
    assert detail["A2"].value == "SUREC-001\nSUREC-002"
    assert detail["B2"].value == 87593849
    assert detail["C5"].value == 3
    assert detail.cell(5, detail.max_column - 3).value == "Kalan eklenti"
    assert detail.cell(5, detail.max_column).value == "Hurdalanmak istenmemiş / eksik eklenti"
    assert detail.max_column == 25
    assert detail.cell(1, detail.max_column - 1).value == "Kısmi Çıkış Oranı (txtKismiCikisOrani)"
    workbook.close()

    print(f"TEST BAŞARILI — 10.004 Amortisman satırı {elapsed:.2f} saniyede tarandı")
    print(output_path)


if __name__ == "__main__":
    run_test()
